import re
from io import BytesIO
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from pypdf import PdfReader

from app.services.jev_service import JEV_MODEL, detect_metadata_batch


BASE_DIR = Path(__file__).resolve().parents[3]
load_dotenv(BASE_DIR / ".env")

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_FILE_COUNT = 6
MAX_EXTRACTED_CHARACTERS = 40_000
MAX_METADATA_CANDIDATES = 64
METADATA_FALLBACKS = {
    "brand": "No brand found",
    "market": "No market found",
    "asset_type": "No asset type found",
}
METADATA_DEFAULTS = {
    "brand": "No brand found",
    "market": "Canada",
    "asset_type": "Email",
}
BRAND_FILENAME_STOP_WORDS = {
    "actual",
    "asset",
    "brief",
    "campaign",
    "copy",
    "dataset",
    "final",
    "guardrail",
    "guardrails",
    "material",
    "reference",
    "relevance",
    "rules",
    "version",
}
MARKETS = {
    "United States": ("United States", "United States of America", "USA", "U.S.A.", "US", "U.S."),
    "United Kingdom": ("United Kingdom", "UK", "U.K.", "Britain"),
    "Canada": ("Canada",),
    "Australia": ("Australia",),
    "New Zealand": ("New Zealand",),
    "India": ("India",),
    "Germany": ("Germany",),
    "France": ("France",),
    "Italy": ("Italy",),
    "Spain": ("Spain",),
    "Japan": ("Japan",),
    "China": ("China",),
    "Brazil": ("Brazil",),
    "Mexico": ("Mexico",),
    "South Africa": ("South Africa",),
    "Singapore": ("Singapore",),
    "Switzerland": ("Switzerland",),
    "Netherlands": ("Netherlands",),
    "Belgium": ("Belgium",),
    "Ireland": ("Ireland",),
    "Sweden": ("Sweden",),
    "Norway": ("Norway",),
    "Denmark": ("Denmark",),
    "Finland": ("Finland",),
    "Poland": ("Poland",),
    "Austria": ("Austria",),
    "Portugal": ("Portugal",),
    "Greece": ("Greece",),
    "Turkey": ("Turkey", "Türkiye"),
    "South Korea": ("South Korea", "Korea"),
    "Taiwan": ("Taiwan",),
    "Hong Kong": ("Hong Kong",),
    "United Arab Emirates": ("United Arab Emirates", "UAE"),
    "Saudi Arabia": ("Saudi Arabia",),
    "Israel": ("Israel",),
    "Argentina": ("Argentina",),
    "Chile": ("Chile",),
    "Colombia": ("Colombia",),
    "Thailand": ("Thailand",),
    "Malaysia": ("Malaysia",),
    "Indonesia": ("Indonesia",),
    "Philippines": ("Philippines",),
    "Vietnam": ("Vietnam", "Viet Nam"),
    "Global": ("Global", "Worldwide", "International"),
    "Europe": ("Europe", "European Union", "EU"),
    "Asia-Pacific": ("Asia-Pacific", "APAC"),
    "Latin America": ("Latin America", "LATAM"),
    "Middle East": ("Middle East", "MENA"),
    "North America": ("North America", "NAM"),
    "Europe, Middle East and Africa": ("Europe, Middle East and Africa", "EMEA"),
}
ASSET_TYPES = (
    "email invitation",
    "email",
    "display ad",
    "banner",
    "brochure",
    "website",
    "landing page",
    "social media post",
    "social media",
    "video",
    "print ad",
    "sales aid",
    "e-detail",
    "e-detailing",
    "flyer",
    "poster",
    "infographic",
    "newsletter",
    "sms",
    "webinar",
    "presentation",
    "leaflet",
    "mobile app",
    "press release",
    "podcast",
    "product monograph",
    "leave-behind",
    "detail aid",
)
ASSET_TYPE_TERMS = {
    term.casefold() for asset_type in ASSET_TYPES for term in asset_type.split()
}
ASSET_TYPE_BY_CASEFOLD = {
    asset_type.casefold(): asset_type.title() for asset_type in ASSET_TYPES
}
NON_BRAND_CAPITALIZED_WORDS = {
    "a", "an", "and", "as", "at", "by", "for", "from", "in", "into",
    "is", "it", "of", "on", "or", "the", "to", "with", "our", "your",
    "this", "that", "these", "those", "project", "brief", "actual",
    "material", "reference", "brand", "market", "asset", "type",
    "campaign", "creative", "copy", "content", "document", "version",
}
BRAND_PHRASE_EDGE_WORDS = {
    "a", "an", "and", "at", "by", "for", "from", "in", "of", "on",
    "the", "to", "with", "campaign", "brand", "product", "asset",
}


def extract_document_text(filename: str, contents: bytes) -> str:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "pdf":
        reader = PdfReader(BytesIO(contents))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if extension == "docx":
        document = Document(BytesIO(contents))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join([*paragraphs, *tables]).strip()
    raise ValueError(
        f"Unsupported file type for {filename}. Use PDF or DOCX."
    )


def _add_candidate(
    candidates: dict[str, list[str]],
    value: str,
    source_label: str,
) -> None:
    normalized = value.strip().strip(" .,:;|-\t")
    if not normalized or len(normalized) > 80:
        return
    existing = next(
        (
            candidate for candidate in candidates
            if candidate.casefold() == normalized.casefold()
        ),
        None,
    )
    key = existing or normalized
    if key not in candidates:
        candidates[key] = []
    if source_label not in candidates[key]:
        candidates[key].append(source_label)


def _field_candidates(
    sources: list[dict[str, str]],
    fields: set[str] | None = None,
) -> dict[str, dict[str, list[str]]]:
    requested_fields = fields if fields is not None else set(METADATA_FALLBACKS)
    candidates: dict[str, dict[str, list[str]]] = {
        field: {} for field in requested_fields
    }
    label_patterns = {
        "brand": (
            r"\b(?:brand(?:\s+name)?|product(?:\s+name)?)"
            r"\s*[:=-]\s*([^\n|;,.]{1,80})"
        ),
        "market": (
            r"\b(?:market|country|region)\s*[:=-]\s*([^\n|;,.]{1,80})"
        ),
        "asset_type": (
            r"\b(?:asset(?:\s+type)?|content\s+type|format|channel)"
            r"\s*[:=-]\s*([^\n|;,.]{1,80})"
        ),
    }
    proper_name_pattern = re.compile(
        r"\b[A-Z][A-Za-z0-9&'-]*(?:\s+[A-Z][A-Za-z0-9&'-]*){0,3}\b"
    )

    for source in sources:
        label, text = source["label"], source["text"]
        if label.casefold().endswith(".xlsx"):
            filename = re.split(r"[/\\]", label)[-1]
            stem = filename.rsplit(".", 1)[0]
            filename_terms = re.split(r"[_\-\s]+", stem)
            brand_terms: list[str] = []
            for term in filename_terms:
                normalized_term = term.casefold()
                is_market = any(
                    normalized_term == alias.casefold()
                    for aliases in MARKETS.values()
                    for alias in aliases
                )
                if (
                    normalized_term in BRAND_FILENAME_STOP_WORDS
                    or normalized_term in ASSET_TYPE_TERMS
                    or is_market
                ):
                    break
                if term:
                    brand_terms.append(term)
            if brand_terms and "brand" in requested_fields:
                _add_candidate(
                    candidates["brand"],
                    " ".join(brand_terms),
                    label,
                )
            continue

        for field, pattern in label_patterns.items():
            if field not in requested_fields:
                continue
            for match in re.finditer(pattern, text, flags=re.IGNORECASE):
                _add_candidate(candidates[field], match.group(1), label)

        for canonical, aliases in MARKETS.items():
            if "market" not in requested_fields:
                break
            for alias in sorted(aliases, key=len, reverse=True):
                if re.search(
                    rf"(?<!\w){re.escape(alias)}(?!\w)",
                    text,
                    re.IGNORECASE,
                ):
                    _add_candidate(candidates["market"], canonical, label)
                    break

        for asset_type in ASSET_TYPES:
            if "asset_type" not in requested_fields:
                break
            if re.search(
                rf"(?<!\w){re.escape(asset_type)}(?!\w)",
                text,
                re.IGNORECASE,
            ):
                _add_candidate(
                    candidates["asset_type"],
                    asset_type.title(),
                    label,
                )

        if "brand" not in requested_fields:
            continue
        for match in proper_name_pattern.finditer(text):
            words = match.group(0).split()
            while words and words[0].casefold() in BRAND_PHRASE_EDGE_WORDS:
                words.pop(0)
            while words and words[-1].casefold() in BRAND_PHRASE_EDGE_WORDS:
                words.pop()
            candidate = " ".join(words)
            if not candidate:
                continue
            if all(
                word.casefold() in NON_BRAND_CAPITALIZED_WORDS
                for word in candidate.split()
            ):
                continue
            if any(
                candidate.casefold() == market.casefold()
                or candidate.casefold()
                in {alias.casefold() for alias in aliases}
                for market, aliases in MARKETS.items()
            ):
                continue
            if any(
                word.casefold() in ASSET_TYPE_TERMS
                for word in candidate.split()
            ):
                continue
            _add_candidate(candidates["brand"], candidate, label)

    too_many = {
        field: len(values)
        for field, values in candidates.items()
        if len(values) > MAX_METADATA_CANDIDATES
    }
    if too_many:
        details = ", ".join(
            f"{field}: {count}" for field, count in too_many.items()
        )
        raise ValueError(
            "Too many candidate values were found for one JEV metadata batch "
            f"({details}; maximum {MAX_METADATA_CANDIDATES} per field). "
            "Narrow the provided material and try again."
        )
    return candidates


def _normalize_explicit_value(field: str, value: str) -> str:
    normalized = value.strip().strip(" .,:;|-\t")
    if field == "market":
        for canonical, aliases in MARKETS.items():
            if normalized.casefold() == canonical.casefold() or any(
                normalized.casefold() == alias.casefold() for alias in aliases
            ):
                return canonical
    if field == "asset_type":
        return ASSET_TYPE_BY_CASEFOLD.get(normalized.casefold(), normalized)
    return normalized


def _explicit_metadata(
    sources: list[dict[str, str]],
) -> dict[str, dict[str, str]]:
    patterns = {
        "brand": (
            r"\b(?:brand(?:\s+name)?|product(?:\s+name)?)\s*[:=-]\s*"
            r"([^\n|;,.]{1,80})"
        ),
        "market": r"\b(?:market|country|region)\s*[:=-]\s*([^\n|;,.]{1,80})",
        "asset_type": (
            r"\b(?:asset(?:\s+type)?|content\s+type|format|channel)"
            r"\s*[:=-]\s*([^\n|;,.]{1,80})"
        ),
    }
    found: dict[str, dict[str, str]] = {}
    for field, pattern in patterns.items():
        for source in sources:
            label = source["label"].casefold()
            if label.endswith(".xlsx"):
                continue
            match = re.search(pattern, source["text"], flags=re.IGNORECASE)
            if match:
                value = _normalize_explicit_value(field, match.group(1))
                if value:
                    found[field] = {
                        "value": value,
                        "source": source["label"],
                    }
                    break
    return found


def _compact_sources(
    sources: list[dict[str, str]],
    candidates: dict[str, dict[str, list[str]]],
) -> list[dict[str, str]]:
    snippets: list[dict[str, str]] = []
    for source in sources:
        text = source["text"]
        excerpts: list[str] = []
        for field_candidates in candidates.values():
            for value, source_labels in field_candidates.items():
                if source["label"] not in source_labels:
                    continue
                match = re.search(re.escape(value), text, flags=re.IGNORECASE)
                if match:
                    start = max(0, match.start() - 160)
                    end = min(len(text), match.end() + 160)
                    excerpt = text[start:end].strip()
                    if excerpt and excerpt not in excerpts:
                        excerpts.append(excerpt)
        if not excerpts and not snippets:
            excerpts.append(text[:320])
        if excerpts:
            snippets.append({
                "label": source["label"],
                "text": "\n...\n".join(excerpts)[:1200],
            })
    if sum(len(source["text"]) for source in snippets) > 8000:
        remaining = 8000
        bounded: list[dict[str, str]] = []
        for source in snippets:
            excerpt = source["text"][:remaining]
            if excerpt:
                bounded.append({"label": source["label"], "text": excerpt})
                remaining -= len(excerpt)
            if remaining <= 0:
                break
        return bounded
    return snippets


def detect_copy_metadata(sources: list[dict[str, str]]) -> dict[str, object]:
    explicit_metadata = _explicit_metadata(sources)
    metadata = dict(explicit_metadata)
    unresolved_fields = set(METADATA_FALLBACKS) - set(explicit_metadata)
    candidates = _field_candidates(sources, unresolved_fields)
    unresolved_options = {
        field: {
            value: f"Appears in: {', '.join(source_labels)}"
            for value, source_labels in field_candidates.items()
        }
        for field, field_candidates in candidates.items()
        if field not in explicit_metadata
    }
    unresolved_options = {
        field: options
        for field, options in unresolved_options.items()
        if options
    }
    if unresolved_options:
        compact_sources = _compact_sources(sources, candidates)
        batch_result = detect_metadata_batch(compact_sources, unresolved_options)
        answers = batch_result["answers"]
        for field in unresolved_options:
            fallback = METADATA_FALLBACKS[field]
            answer = answers.get(field)
            value = answer.get("choice") if isinstance(answer, dict) else None
            if not isinstance(value, str):
                raise ValueError(
                    f"JEV did not return a valid {field.replace('_', ' ')} choice."
                )
            if value == fallback:
                default_value = METADATA_DEFAULTS[field]
                metadata[field] = {
                    "value": default_value,
                    "source": (
                        "Default value" if default_value != fallback else "No source found"
                    ),
                }
                continue
            source_labels = candidates[field].get(value)
            if not source_labels:
                raise ValueError(
                    f"JEV returned an unknown {field.replace('_', ' ')} choice."
                )
            metadata[field] = {"value": value, "source": source_labels[0]}

    for field, default_value in METADATA_DEFAULTS.items():
        metadata.setdefault(
            field,
            {
                "value": default_value,
                "source": (
                    "Default value"
                    if default_value != METADATA_FALLBACKS[field]
                    else "No source found"
                ),
            },
        )

    classification_context = "\n\n".join(
        f"{source['label']}:\n{source['text']}" for source in sources
    )
    return {
        "metadata": metadata,
        "model": JEV_MODEL if unresolved_options else "direct labeled values",
        "classification_context": classification_context,
        "warnings": [
            f"{source['label']}: only the spreadsheet filename is included; its contents are not analyzed."
            for source in sources
            if source["label"].casefold().endswith(".xlsx")
        ],
    }
