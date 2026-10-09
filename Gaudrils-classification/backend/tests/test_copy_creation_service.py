import unittest
from io import BytesIO
from unittest.mock import Mock, patch

from docx import Document

from app.services.copy_creation_service import (
    _field_candidates,
    detect_copy_metadata,
    extract_document_text,
)
from app.services.jev_service import JEV_MODEL


class DocumentTextExtractionTests(unittest.TestCase):
    def test_extracts_docx_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_paragraph("Jardiance in Canada")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Asset"
        table.cell(0, 1).text = "Email"
        contents = BytesIO()
        document.save(contents)

        text = extract_document_text("brief.docx", contents.getvalue())

        self.assertIn("Jardiance in Canada", text)
        self.assertIn("Asset | Email", text)

    def test_rejects_unsupported_file_types(self) -> None:
        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            extract_document_text("brief.txt", b"contents")

    def test_spreadsheet_content_is_not_extracted(self) -> None:
        with self.assertRaisesRegex(ValueError, "Use PDF or DOCX"):
            extract_document_text("large.xlsx", b"spreadsheet contents")


class CopyMetadataDetectionTests(unittest.TestCase):
    @patch("app.services.copy_creation_service.detect_metadata_batch")
    def test_labeled_descriptions_are_returned_without_waiting_for_jev(
        self,
        detect_batch: Mock,
    ) -> None:
        result = detect_copy_metadata([
            {
                "label": "Project Brief description",
                "text": "brand: mankind, asset: email",
            },
            {
                "label": "Actual Material description",
                "text": "market : india",
            },
        ])

        detect_batch.assert_not_called()
        self.assertEqual(
            result["metadata"],
            {
                "brand": {
                    "value": "mankind",
                    "source": "Project Brief description",
                },
                "market": {"value": "India", "source": "Actual Material description"},
                "asset_type": {
                    "value": "Email",
                    "source": "Project Brief description",
                },
            },
        )

    def test_builds_supported_brand_market_and_asset_candidates(self) -> None:
        candidates = _field_candidates([
            {
                "label": "Actual Material / Canada_Email.docx",
                "text": (
                    "Jardiance email invitation for Canadian healthcare "
                    "professionals. Market: Canada."
                ),
            }
        ])

        self.assertIn("Jardiance", candidates["brand"])
        self.assertIn("Canada", candidates["market"])
        self.assertIn("Email Invitation", candidates["asset_type"])
        self.assertEqual(
            candidates["brand"]["Jardiance"],
            ["Actual Material / Canada_Email.docx"],
        )

    def test_infers_brand_from_spreadsheet_filename_only(self) -> None:
        candidates = _field_candidates([
            {
                "label": "Project Brief / Jardiance_Guardrails_Relevance.xlsx",
                "text": (
                    "Spreadsheet filename: Jardiance_Guardrails_Relevance.xlsx\n"
                    "This text must not be analyzed."
                ),
            }
        ])

        self.assertEqual(candidates["brand"], {
            "Jardiance": ["Project Brief / Jardiance_Guardrails_Relevance.xlsx"],
        })
        self.assertEqual(candidates["market"], {})
        self.assertEqual(candidates["asset_type"], {})

    @patch("app.services.copy_creation_service.detect_metadata_batch")
    def test_filename_brand_and_default_market_and_asset_values(
        self,
        detect_batch: Mock,
    ) -> None:
        detect_batch.return_value = {
            "answers": {"brand": {"choice": "Jardiance"}},
            "usage": {},
        }

        result = detect_copy_metadata([
            {
                "label": "Project Brief / Jardiance_Guardrails_Relevance.xlsx",
                "text": (
                    "Spreadsheet filename: Jardiance_Guardrails_Relevance.xlsx"
                ),
            }
        ])

        self.assertEqual(
            detect_batch.call_args.args[1],
            {
                "brand": {
                    "Jardiance": (
                        "Appears in: Project Brief / "
                        "Jardiance_Guardrails_Relevance.xlsx"
                    )
                }
            },
        )
        self.assertEqual(
            result["metadata"],
            {
                "brand": {
                    "value": "Jardiance",
                    "source": "Project Brief / Jardiance_Guardrails_Relevance.xlsx",
                },
                "market": {"value": "Canada", "source": "Default value"},
                "asset_type": {"value": "Email", "source": "Default value"},
            },
        )

    @patch("app.services.copy_creation_service.detect_metadata_batch")
    def test_uses_one_jev_batch_and_keeps_source_provenance(
        self,
        detect_batch: Mock,
    ) -> None:
        detect_batch.return_value = {
            "answers": {
                "brand": {"choice": "Jardiance"},
                "market": {"choice": "Canada"},
                "asset_type": {"choice": "Email Invitation"},
            },
            "usage": {},
        }
        sources = [
            {
                "label": "Project Brief description",
                "text": "Canada market. Jardiance campaign.",
            },
            {
                "label": "Actual Material / campaign.docx",
                "text": "Jardiance email invitation.",
            },
        ]

        result = detect_copy_metadata(sources)

        detect_batch.assert_called_once()
        self.assertEqual(detect_batch.call_args.args[0], sources)
        self.assertEqual(result["model"], JEV_MODEL)
        self.assertEqual(
            result["metadata"]["brand"],
            {"value": "Jardiance", "source": "Project Brief description"},
        )
        self.assertEqual(
            result["metadata"]["market"],
            {"value": "Canada", "source": "Project Brief description"},
        )
        self.assertEqual(
            result["metadata"]["asset_type"],
            {"value": "Email Invitation", "source": "Actual Material / campaign.docx"},
        )
        self.assertIn(
            "Project Brief description:\nCanada market. Jardiance campaign.",
            result["classification_context"],
        )
        options = detect_batch.call_args.args[1]
        self.assertIn("Jardiance", options["brand"])
        self.assertIn("Canada", options["market"])
        self.assertIn("Email Invitation", options["asset_type"])

    @patch("app.services.copy_creation_service.detect_metadata_batch")
    def test_uses_defaults_when_jev_has_no_market_or_asset_match(
        self,
        detect_batch: Mock,
    ) -> None:
        detect_batch.return_value = {
            "answers": {
                "brand": {"choice": "No brand found"},
                "market": {"choice": "No market found"},
                "asset_type": {"choice": "No asset type found"},
            },
            "usage": {},
        }

        result = detect_copy_metadata([
            {"label": "Actual Material description", "text": "Generic copy"}
        ])

        self.assertEqual(
            result["metadata"],
            {
                "brand": {"value": "No brand found", "source": "No source found"},
                "market": {"value": "Canada", "source": "Default value"},
                "asset_type": {
                    "value": "Email",
                    "source": "Default value",
                },
            },
        )

    @patch("app.services.copy_creation_service.detect_metadata_batch")
    def test_uses_defaults_when_no_candidates_exist(
        self,
        detect_batch: Mock,
    ) -> None:
        result = detect_copy_metadata([
            {
                "label": "Actual Material description",
                "text": "generic copy without names",
            }
        ])

        detect_batch.assert_not_called()
        self.assertEqual(
            result["metadata"]["market"],
            {"value": "Canada", "source": "Default value"},
        )
        self.assertEqual(
            result["metadata"]["asset_type"],
            {"value": "Email", "source": "Default value"},
        )


if __name__ == "__main__":
    unittest.main()
